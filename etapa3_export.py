"""
ETAPA 3 — Exportação final (CSV, XLSX formatado, metadados e diff mensal)
"""

import os
import glob
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import OUTPUT_DIR, HISTORICO_DIR
from logger_setup import log

COLUNAS_FINAIS = [
    "score_lead", "cnpj", "razao_social", "nome_fantasia", "porte_empresa", "porte_real",
    "eh_matriz", "cnae_fiscal_principal", "cnae_descricao", "situacao_cadastral",
    "data_abertura",                                   # NOVO — requisito 2
    "ddd1", "telefone1", "ddd2", "telefone2",
    "telefone_celular", "possui_celular",              # NOVO — requisito 1
    "email", "email_confiavel", "email_confiavel_motivo",
    "logradouro", "numero", "bairro", "municipio_nome", "uf", "cep",
    "fonte", "nome_encontrado", "score_similaridade", "site_institucional", "situacao_mec",
]

def preparar_dataframe_final(df: pd.DataFrame) -> pd.DataFrame:
    for col in COLUNAS_FINAIS:
        if col not in df.columns:
            df[col] = None
    return df[COLUNAS_FINAIS].copy()

def exportar_csv(df: pd.DataFrame, nome_arquivo: str) -> str:
    caminho = os.path.join(OUTPUT_DIR, nome_arquivo)
    df.to_csv(caminho, index=False, encoding="utf-8-sig")
    log.info(f"CSV exportado: {caminho}")
    return caminho

def salvar_snapshot_historico(df: pd.DataFrame) -> str:
    timestamp = datetime.now().strftime("%Y-%m")
    caminho = os.path.join(HISTORICO_DIR, f"snapshot_{timestamp}.parquet")
    df.to_parquet(caminho, index=False)
    return caminho

def carregar_snapshot_anterior(excluir_timestamp_atual: str) -> pd.DataFrame | None:
    snapshots = sorted(glob.glob(os.path.join(HISTORICO_DIR, "snapshot_*.parquet")))
    snapshots = [s for s in snapshots if excluir_timestamp_atual not in s]
    if not snapshots:
        return None
    return pd.read_parquet(snapshots[-1])

def calcular_novidades(df_atual: pd.DataFrame) -> pd.DataFrame:
    timestamp_atual = datetime.now().strftime("%Y-%m")
    df_anterior = carregar_snapshot_anterior(timestamp_atual)

    if df_anterior is None or df_anterior.empty:
        return pd.DataFrame(columns=df_atual.columns)

    cnpjs_anteriores = set(df_anterior["cnpj"])
    novidades = df_atual[~df_atual["cnpj"].isin(cnpjs_anteriores)].copy()
    log.info(f"{len(novidades)} empresas novas em relação ao snapshot anterior.")
    return novidades

def _formatar_aba(ws, df: pd.DataFrame):
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    for col_idx in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    corpo_font = Font(name="Arial", size=10)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = corpo_font

    for col_idx, col_nome in enumerate(df.columns, start=1):
        maior_valor = df[col_nome].fillna("").astype(str).str.len().max() if len(df) else 10
        largura = min(max(len(str(col_nome)), int(maior_valor)) + 2, 45)
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions

def montar_metadados(df: pd.DataFrame, cnaes_alvo: dict) -> pd.DataFrame:
    municipios_encontrados = ", ".join(df["municipio_nome"].dropna().unique()) if "municipio_nome" in df.columns else "N/A"
    linhas = [
        ("Data de extração", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Município(s) processado(s)", municipios_encontrados),
        ("Total de empresas exportadas", len(df)),
    ]
    for cnae, descricao in cnaes_alvo.items():
        qtd = int((df["cnae_fiscal_principal"] == cnae).sum()) if "cnae_fiscal_principal" in df and len(df) else 0
        linhas.append((f"  CNAE {cnae}", qtd))

    return pd.DataFrame(linhas, columns=["Métrica", "Valor"])

def exportar_xlsx(df, cnaes_alvo, novidades=None, nome_arquivo="") -> str:
    caminho = os.path.join(OUTPUT_DIR, nome_arquivo)

    # Requisito 1: só quem tem celular, já no formato pedido
    df_celulares = df[df["possui_celular"] == True][["razao_social", "telefone_celular"]].copy()
    df_celulares = df_celulares.rename(columns={"telefone_celular": "celular"})

    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        montar_metadados(df, cnaes_alvo).to_excel(writer, sheet_name="Metadados", index=False)
        df.to_excel(writer, sheet_name="Leads B2B", index=False)
        if novidades is not None and not novidades.empty:
            novidades.to_excel(writer, sheet_name="Novidades", index=False)
        if not df_celulares.empty:
            df_celulares.to_excel(writer, sheet_name="Somente Celular", index=False)

    wb = load_workbook(caminho)
    _formatar_aba(wb["Metadados"], montar_metadados(df, cnaes_alvo))
    _formatar_aba(wb["Leads B2B"], df)
    if "Novidades" in wb.sheetnames:
        _formatar_aba(wb["Novidades"], novidades)
    if "Somente Celular" in wb.sheetnames:
        _formatar_aba(wb["Somente Celular"], df_celulares)

    wb.save(caminho)
    log.info(f"XLSX exportado: {caminho}")
    return caminho

def rodar_etapa3(df: pd.DataFrame, cnaes_alvo: dict) -> tuple[str, str]:
    log.info("=== ETAPA 3: Exportação ===")

    # REQUISITO 3: Filtragem estrita - Garante que APENAS os CNAEs digitados vão para o Excel
    df_filtrado = df[df["cnae_fiscal_principal"].isin(cnaes_alvo.keys())].copy()

    df_final = preparar_dataframe_final(df_filtrado)
    novidades = calcular_novidades(df_final)
    salvar_snapshot_historico(df_final)

    # REQUISITO 1: Nome do arquivo baseado no CNAE e na Data
    data_hoje = datetime.now().strftime("%Y-%m-%d")
    lista_cnaes = list(cnaes_alvo.keys())

    # Pega os 2 primeiros CNAEs para não deixar o nome do arquivo gigante
    str_cnaes = "_".join(lista_cnaes[:2])
    if len(lista_cnaes) > 2:
        str_cnaes += "_etc"

    nome_base = f"leads_cnae_{str_cnaes}_{data_hoje}"

    caminho_csv = exportar_csv(df_final, f"{nome_base}.csv")
    caminho_xlsx = exportar_xlsx(df_final, cnaes_alvo, novidades, f"{nome_base}.xlsx")

    return caminho_csv, caminho_xlsx